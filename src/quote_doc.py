"""The one customer-facing quotation model. The PDF, the preview JSON and the XLSX / CSV exports all come from compose(), so they
cannot disagree. Selling prices only: cost and margin enter only through the internal export (costs() + internal=True)."""
import csv
import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .pricing import _q
from .safe import safe

DEFAULT_COLOR = "#1f6f4a"


def _trim(q) -> Decimal:
    s = f"{q:f}"
    return Decimal(s.rstrip("0").rstrip(".") if "." in s else s)


def compose(org: dict, doc: dict, customer: str, lines: list[dict], party: dict | None = None) -> dict:
    """org / doc / lines are the rows quote_pdf.render always took (extra org keys are optional). -> the model dict."""
    items, sub, tax, taxed = [], Decimal(0), Decimal(0), False
    for l in lines:
        if l["is_header"]:
            items.append({"no": None, "description": l["description"], "qty": None, "uom": None, "unit_price": None,
                          "discount_pct": None, "amount": None, "is_header": True, "section": l.get("section")})
            continue
        # a catalog line can still await its price (unit_price NULL): it counts as 0 until staff set it
        amt = _q(l["qty"] * (l["unit_price"] or 0) * (1 - (l["discount_pct"] or 0) / Decimal(100)))
        sub += amt
        if org["sst_registered"] and l.get("tax_rate") is not None:  # never invent a rate: no per-line rate, no SST line
            tax, taxed = tax + _q(amt * l["tax_rate"] / 100), True
        items.append({"no": l["line_no"], "description": l["description"], "qty": _trim(l["qty"]), "uom": l["uom"],
                      "unit_price": l["unit_price"], "discount_pct": l["discount_pct"], "amount": amt,
                      "is_header": False, "section": l.get("section")})
    on = doc["issued_on"] or doc["created_at"]
    p = party or {}
    return {"doc_no": doc["doc_no"], "status": doc.get("status"),
            "customer": {"name": customer, "attention": p.get("attention"), "address": p.get("address"), "email": p.get("email")},
            "issued_on": on.date() if isinstance(on, datetime) else on, "valid_until": doc["valid_until"], "currency": doc["currency"],
            "lines": items, "subtotal": sub, "tax_total": tax if taxed else None, "total": sub + tax,
            "payment_terms": doc.get("payment_terms") or org.get("quote_payment_terms"),  # the quote's own, else the business default
            "terms_text": doc.get("terms_text") or org.get("quote_terms_text"),
            "issuer": {"name": org["name"], "address": org["address"], "ssm_no": org["ssm_no"],
                       "sst_no": org["sst_no"] if org["sst_registered"] else None, "email": org.get("issuer_email"),
                       "phone": org.get("issuer_phone"), "logo_url": org.get("logo_data_url"),
                       "brand_color": org.get("brand_color") or DEFAULT_COLOR, "bank_details": org["bank_details"]}}


def load(conn, doc_id) -> dict:
    """compose() from the database. The caller has already checked the document is this tenant's out quote."""
    d = conn.execute("""select d.doc_no, d.status, d.currency, d.issued_on, d.created_at, d.valid_until, d.payment_terms, d.terms_text,
                               p.name as party_name, p.attention, p.address, p.email, d.tenant_id
                        from documents d join parties p on p.id = d.party_id where d.id=%s""", (doc_id,)).fetchone()
    lines = conn.execute("select * from document_lines where document_id=%s order by line_no", (doc_id,)).fetchall()
    return compose(org(conn, d["tenant_id"]), d, d["party_name"], lines, d)


def org(conn, tenant_id) -> dict:
    return conn.execute(
        """select coalesce(display_name, name) as name, ssm_no, address, bank_details, sst_registered, sst_no, logo_data_url, brand_color,
                  issuer_email, issuer_phone, quote_payment_terms, quote_terms_text, currency from tenants where id=%s""", (tenant_id,)).fetchone()


def costs(conn, doc_id) -> dict:
    """INTERNAL: line_no -> {cost, stale} from the margin view (no row for header or cost-less lines)."""
    return {x["line_no"]: {"cost": _q(x["line_cost"]), "stale": x["stale"]} for x in conn.execute(
        """select l.line_no, m.line_cost, coalesce(bool_or(p.cost_is_stale), false) as stale
           from v_quote_margin m join document_lines l on l.id = m.line_id
           left join v_quote_provenance p on p.line_id = m.line_id
           where m.quote_id=%s group by l.line_no, m.line_cost""", (doc_id,))}


def _sheet(m: dict, cost: dict | None) -> list[list]:
    """Every row of the XLSX / CSV: header block, table, totals, terms. Numbers stay numbers (Decimal); text goes through safe()."""
    i, c = m["issuer"], m["customer"]
    cur, t = m["currency"], lambda *r: [safe(v) for v in r]
    rows = [t("Quotation", m["doc_no"]), t("Issuer", i["name"]), t("Address", i["address"]), t("SSM no", i["ssm_no"]), t("SST no", i["sst_no"]),
            t("Email", i["email"]), t("Phone", i["phone"]), [], t("Customer", c["name"]), t("Attention", c["attention"]),
            t("Customer address", c["address"]), t("Customer email", c["email"]), t("Issued", str(m["issued_on"])),
            t("Valid until", str(m["valid_until"]) if m["valid_until"] else None), t("Currency", cur), []]
    head = ["#", "Description", "Qty", "Unit", f"Unit price ({cur})", "Discount %", f"Amount ({cur})"]
    if cost is not None:
        head += [f"Cost ({cur})", f"Margin ({cur})", "Margin %", "Cost stale"]
    rows.append(head)
    tc = Decimal(0)
    for l in m["lines"]:
        if l["is_header"]:
            rows.append(["", *t(l["description"])])
            continue
        row = [l["no"], *t(l["description"]), l["qty"], *t(l["uom"]), l["unit_price"], l["discount_pct"], l["amount"]]
        if cost is not None:
            x = cost.get(l["no"])
            if x:
                tc += x["cost"]
                mg = l["amount"] - x["cost"]
                row += [x["cost"], mg, round(mg / l["amount"] * 100, 2) if l["amount"] else None, "yes" if x["stale"] else ""]
            else:
                row += [None, None, None, ""]
        rows.append(row)
    pad = [""] * 4
    rows.append([])
    rows.append([*pad, "", "Subtotal", m["subtotal"]])
    if m["tax_total"] is not None:
        rows.append([*pad, "", "SST", m["tax_total"]])
    rows.append([*pad, "", "Total", m["total"]])
    if cost is not None:
        mg = m["subtotal"] - tc
        rows += [[], t("Total cost", tc), t("Total margin", mg), t("Margin %", round(mg / m["subtotal"] * 100, 2) if m["subtotal"] else None),
                 t("Any cost stale", "yes" if any(x["stale"] for x in cost.values()) else "no")]
    rows += [[], t("Payment terms", m["payment_terms"]), t("Terms", m["terms_text"]), t("Bank details", i["bank_details"])]
    return [r for r in rows if not (r and r[-1] is None and len(r) == 2)]  # drop "Label, (nothing)" rows


def xlsx(m: dict, cost: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    for r in _sheet(m, cost):
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    for n, w in enumerate((6, 48, 10, 8, 18, 12, 18, 16, 16, 10, 11), 1):
        ws.column_dimensions[get_column_letter(n)].width = w
    for row in ws.iter_rows(min_col=1, max_col=1):
        if isinstance(row[0].value, str) and row[0].value in ("Quotation", "#", "Total", "Total cost"):
            row[0].font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(m: dict, cost: dict | None = None) -> bytes:
    """utf-8 with a BOM so Excel reads it as UTF-8."""
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in _sheet(m, cost):
        w.writerow([f"{v:f}" if isinstance(v, Decimal) else v for v in r])
    return buf.getvalue().encode("utf-8-sig")
