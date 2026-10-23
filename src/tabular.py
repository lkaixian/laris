"""XLSX / DOCX / CSV -> plain text rows for the extraction model. Deterministic, no OCR. Raises ValueError on a bad file
(the pipeline turns that into a 422); legacy .xls/.doc are never sniffed, so they end up as a 415 there."""
import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from xml.etree import ElementTree as ET

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAGIC = ((b"%PDF", "application/pdf"), (b"\xff\xd8\xff", "image/jpeg"), (b"\x89PNG\r\n\x1a\n", "image/png"),
         (b"RIFF", "image/webp"), (b"PK\x03\x04", "zip"))  # the client-sent content type is not trusted: sniff
MAX_UNZIPPED = 50_000_000
MAX_ROWS, MAX_COLS = 2000, 40
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:  # ponytail: gb18030 when every non-ASCII char is CJK; a Latin-1 file that happens to pass is the ceiling
        s = data.decode("gb18030")
        if sum("一" <= c <= "鿿" for c in s) >= 2 and all(c < "\x80" or "一" <= c <= "鿿" for c in s):
            return s
    except UnicodeDecodeError:
        pass
    return data.decode("cp1252", errors="replace")


def _csv_rows(data: bytes) -> list[list[str]]:
    s = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(s[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(s), dialect))


def _guard(z: zipfile.ZipFile) -> None:
    if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
        raise ValueError("file expands to more than 50 MB")


def sniff(data: bytes, filename: str | None = None, declared: str | None = None) -> str | None:
    """Real type from the bytes: a mime string, or None if unsupported. A zip is XLSX or DOCX by its member names;
    CSV by a .csv name, or by a text/csv claim whose bytes are text with delimiters."""
    m = next((m for sig, m in MAGIC if data.startswith(sig)), None)
    if m == "zip":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
        except zipfile.BadZipFile:
            return None
        if "[Content_Types].xml" in names:
            if any(n.startswith("xl/") for n in names):
                return XLSX
            if any(n.startswith("word/") for n in names):
                return DOCX
        return None
    if m or data.startswith(b"\xd0\xcf\x11\xe0") or b"\x00" in data[:1024]:  # known type, legacy OLE .xls/.doc, binary
        return m
    if (filename or "").lower().endswith(".csv"):
        return "text/csv"
    if (declared or "").split(";")[0].lower() in ("text/csv", "application/csv"):
        head = _decode(data[:4096]).split("\n", 1)[0]
        return "text/csv" if any(d in head for d in ",;\t") else None
    return None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.10g}"  # 0.1+0.2 -> 0.3, not 0.30000000000000004
    if isinstance(v, datetime):  # Excel stores every date as a datetime
        return (v.date() if v.time() == time() else v).isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _capped(sheets: list[tuple[str | None, list]]) -> str:
    out, n, cut = [], 0, False
    for name, rows in sheets:
        if name is not None:
            out.append(f"## sheet: {name}")
        for row in rows:
            cells = [_cell(c) for c in row]
            if not any(cells):
                continue
            if n >= MAX_ROWS:
                cut = True
                break
            cut |= len(cells) > MAX_COLS
            out.append("\t".join(cells[:MAX_COLS]))
            n += 1
    if cut:
        out.append(f"[truncated: only the first {MAX_ROWS} rows and {MAX_COLS} columns were read]")
    return "\n".join(out)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        _guard(z)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return _capped([(ws.title, ws.iter_rows(values_only=True)) for ws in wb.worksheets])
    finally:
        wb.close()


def _text(el) -> str:
    return "".join("\t" if t.tag == W + "tab" else t.text or "" for t in el.iter() if t.tag in (W + "t", W + "tab"))


def _docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        _guard(z)
        xml = z.read("word/document.xml")
    if re.search(rb"<!(DOCTYPE|ENTITY)", xml.replace(b"\x00", b""), re.I):  # XXE / billion laughs: never parse
        raise ValueError("document contains a DOCTYPE/ENTITY declaration")
    body = ET.fromstring(xml).find(W + "body")
    if body is None:
        raise ValueError("not a Word document")
    lines = []
    for el in body:
        if el.tag == W + "p":
            lines.append(_text(el))
        elif el.tag == W + "tbl":
            lines += ["\t".join(" ".join(_text(p) for p in tc.iter(W + "p")) for tc in tr.findall(W + "tc"))
                      for tr in el.iter(W + "tr")]
    return "\n".join(ln for ln in lines if ln.strip())


def to_text(data: bytes, mime: str) -> str:
    if mime == XLSX:
        return _xlsx(data)
    if mime == DOCX:
        return _docx(data)
    return _capped([(None, _csv_rows(data))])
