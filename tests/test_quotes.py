"""Step 106: quotation defaults, the one document model (PDF = preview = exports), exports, and sending the PDF. Nothing touches a network."""
import base64
import csv
import io
import json
from datetime import timedelta
from decimal import Decimal

import httpx
import psycopg
import pytest
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from src import db, mailer, pipeline, quote_pdf
from src import license as lic
from src.api import deps
from src.api.deps import STAFF_HIDE, Services
from src.channels import telegram, whatsapp
from support.fixtures import env, pg_url  # noqa: F401
from test_api import keys, seed_quote, tr
from test_license import PUB, Net
from test_roles import member, role
from test_rfq import Meta, other_tenant

# the seeded quote: 40 x RM60 = RM2,400 selling; cost 1,500 and margin 900 (37.5 %) must never reach the customer
COST_WORDS = ("1500", "1,500", "900.00", "37.5", "margin", "cost", "supplier")


def pdf_text(b: bytes) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(b)).pages)


def xlsx_rows(b: bytes):
    return list(load_workbook(io.BytesIO(b)).active.iter_rows(values_only=True))


def csv_rows(r):
    assert r.content[:3] == b"\xef\xbb\xbf"  # BOM: Excel opens it as UTF-8
    return list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))


def export(e, q, fmt="pdf", variant="customer", who=None):
    return e.c.get(f"/api/documents/{q}/export?format={fmt}&variant={variant}", headers=who or e.O)


def draft(e, **kw):
    """A fresh draft quote (the cost book is stocked once per test)."""
    if not getattr(e, "stocked", False):
        e.stock_costs()
        e.stocked = True
    r = e.enquire([{"item_key": "print_part_small", "qty": 10}], **kw)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def approved(e, **party):
    """A draft quote the owner approved, its customer given `party` columns (email, contact)."""
    d = draft(e)
    assert tr(e, d, "approved").status_code == 200
    for k, v in party.items():
        e.sql(f"update parties set {k}=%s where id=(select party_id from documents where id=%s)", (v, d))
    return d


def send(e, d, channel="email", who=None, **kw):
    return e.c.post(f"/api/documents/{d}/send", json={"channel": channel, **kw}, headers=who or e.O)


def png(w, h) -> str:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 30, 30)).save(buf, "PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


# ---------------------------------------------------------------- 1. tenant defaults
def test_defaults_feed_new_quotes_and_the_pdf(env):
    e = env
    e.stock_costs()
    first = e.enquire([{"item_key": "print_part_small", "qty": 10}]).json()["id"]
    d0 = e.sql("select doc_no, valid_until, payment_terms, terms_text from documents where id=%s", (first,))[0]
    assert d0["doc_no"].startswith("QT-") and d0["payment_terms"] is None and d0["valid_until"] == pipeline.today_my() + timedelta(days=30)
    r = e.c.post("/api/settings", json={"quote_validity_days": 14, "quote_payment_terms": "50% deposit, balance on delivery",
                                         "quote_terms_text": "Warranty 1 year", "quote_prefix": "SG-A", "issuer_email": "sales@shop.test",
                                         "issuer_phone": "+60 3 1234 5678"}, headers=e.O)
    assert r.status_code == 200 and r.json()["quote_prefix"] == "SG-A" and r.json()["quote_validity_days"] == 14
    d = e.enquire([{"item_key": "print_part_small", "qty": 10}]).json()["id"]
    row = e.sql("select doc_no, valid_until, payment_terms, terms_text from documents where id=%s", (d,))[0]
    year = pipeline.today_my().year
    assert row["doc_no"] == f"SG-A-{year}-0002"  # the counter is per tenant and year, not per prefix
    assert row["valid_until"] == pipeline.today_my() + timedelta(days=14)
    assert (row["payment_terms"], row["terms_text"]) == ("50% deposit, balance on delivery", "Warranty 1 year")
    e.c.post("/api/settings", json={"quote_prefix": "XY"}, headers=e.O)
    e.enquire([{"item_key": "print_part_small", "qty": 10}])
    assert e.sql("select max(doc_no) m from documents where doc_no like %s", (f"XY-{year}-%",))[0]["m"] == f"XY-{year}-0003"
    t = pdf_text(e.c.get(f"/api/documents/{d}/pdf", headers=e.S).content)
    assert "Valid until" in t and "50% deposit" in t and "Warranty 1 year" in t and "sales@shop.test" in t and "+60 3 1234 5678" in t
    e.sql("update documents set payment_terms='Cash on delivery' where id=%s", (d,))  # the quote's own terms win over the default
    t = pdf_text(e.c.get(f"/api/documents/{d}/pdf", headers=e.S).content)
    assert "Cash on delivery" in t and "50% deposit" not in t and "Warranty 1 year" in t  # terms_text still falls back


def test_settings_validation_null_clears_and_audit(env):
    e = env
    bad = [{"quote_validity_days": 0}, {"quote_validity_days": 366}, {"quote_prefix": "qt"}, {"quote_prefix": "TOOLONG123"},
           {"quote_prefix": "A B"}, {"issuer_email": "nope"}, {"issuer_phone": "1" * 41}, {"quote_payment_terms": "x" * 501},
           {"quote_terms_text": "x" * 2001}]
    for b in bad:
        assert e.c.post("/api/settings", json=b, headers=e.O).status_code == 422, b
    assert e.c.post("/api/settings", json={"quote_prefix": "QX"}, headers=e.S).status_code == 403
    ok = {"quote_validity_days": 365, "quote_payment_terms": "x" * 500, "quote_terms_text": "y" * 2000, "quote_prefix": "A1-B2",
          "issuer_email": "a@b.co", "issuer_phone": "1" * 40}
    assert e.c.post("/api/settings", json=ok, headers=e.O).status_code == 200
    got = e.c.get("/api/settings", headers=e.O).json()
    assert {k: got[k] for k in ok} == ok
    cleared = e.c.post("/api/settings", json={k: None for k in ok}, headers=e.O).json()
    assert cleared["quote_validity_days"] == 30 and cleared["quote_prefix"] == "QT"  # NOT NULL columns go back to the default
    assert all(cleared[k] is None for k in ("quote_payment_terms", "quote_terms_text", "issuer_email", "issuer_phone"))
    ev = [x["event"] for x in e.c.get("/api/audit", headers=e.O).json()]
    assert any(x.startswith("settings changed: ") and "quote_prefix" in x for x in ev)
    assert e.c.get("/api/settings", headers=e.S).status_code == 403


# ---------------------------------------------------------------- 2. PDF
def test_pdf_has_page_numbers_and_the_sst_line_only_with_a_rate(env):
    e = env
    q = seed_quote(e)
    assert "Page 1 of 1" in pdf_text(e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content)
    e.sql("update tenants set sst_registered=true, sst_no='W10-1'")
    t = pdf_text(e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content)
    assert "SST W10-1" in t and "Subtotal" in t and "2,400.00" in t and "SST 144" not in t  # registered, but no line rate: no invented SST
    assert e.c.get(f"/api/documents/{q}/preview", headers=e.O).json()["tax_total"] is None
    e.sql("update document_lines set tax_rate=6 where document_id=%s", (q,))
    t = pdf_text(e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content)
    assert "144.00" in t and "2,544.00" in t
    e.sql("update tenants set sst_registered=false")  # not registered: never an SST line, whatever the lines say
    assert e.c.get(f"/api/documents/{q}/preview", headers=e.O).json()["tax_total"] is None
    e.sql("update tenants set sst_registered=true")
    e.sql("""insert into document_lines(document_id, line_no, description, qty, uom, unit_price, price_source)
             select %s, n + 1, 'SAMPLE extra line ' || n, 1, 'pcs', 1, 'fixed_rate_card' from generate_series(1, 70) n""", (q,))
    pdf = e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content
    pages = len(PdfReader(io.BytesIO(pdf)).pages)
    assert pages >= 2 and f"Page {pages} of {pages}" in pdf_text(pdf)


def test_header_colour_is_readable_and_falls_back():
    from reportlab.lib import colors
    assert quote_pdf._header_colors("#ffff00")[1] == colors.black and quote_pdf._header_colors("#102030")[1] == colors.white
    assert quote_pdf._header_colors(None)[0] == colors.HexColor("#1f6f4a") and quote_pdf._header_colors("garbage")[0] == colors.HexColor("#1f6f4a")


def test_logo_good_huge_and_corrupt_never_break_the_pdf(env):
    e = env
    q = seed_quote(e)
    plain = e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content
    assert b"/Subtype /Image" not in plain
    assert e.c.post("/api/settings", json={"logo_data_url": png(400, 100)}, headers=e.O).status_code == 200
    with_logo = e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content
    assert with_logo[:5] == b"%PDF-" and b"/Subtype /Image" in with_logo
    wide, tall = quote_pdf._logo(png(400, 100)), quote_pdf._logo(png(200, 200))
    assert (wide.drawWidth, wide.drawHeight) == (160, 40) and (tall.drawWidth, tall.drawHeight) == (40, 40)  # inside 160 x 40, aspect kept
    assert e.c.get(f"/api/documents/{q}/preview", headers=e.O).json()["issuer"]["logo_url"].startswith("data:image/png;base64,")
    for junk in (png(5000, 5000), "data:image/png;base64,AAAA", "data:image/png;base64,!!!", "http://evil.test/x.png", "data:image/svg+xml;base64,PHN2Zz4="):
        e.sql("update tenants set logo_data_url=%s", (junk,))  # bypasses the API check on purpose: the renderer must cope on its own
        r = e.c.get(f"/api/documents/{q}/pdf", headers=e.O)
        assert r.status_code == 200 and r.content[:5] == b"%PDF-" and b"/Subtype /Image" not in r.content, junk[:30]


# ---------------------------------------------------------------- 3. preview
def test_preview_is_the_pdf_model_and_staff_safe(env):
    e = env
    q = seed_quote(e)
    e.sql("update parties set email='buyer@cafe.test', attention='Ms Lim', address='9 Jalan Kopi' where id=(select party_id from documents where id=%s)", (q,))
    e.sql("update tenants set address='1 Jalan Test', bank_details='Bank 123', issuer_email='me@shop.test', brand_color='#aa0000'")
    p = e.c.get(f"/api/documents/{q}/preview", headers=e.O).json()
    assert p["doc_no"] == "SAMPLE-QT-0001" and p["status"] == "sent" and p["currency"] == "MYR"
    assert p["customer"] == {"name": "SAMPLE Customer Cafe (synthetic)", "attention": "Ms Lim", "address": "9 Jalan Kopi", "email": "buyer@cafe.test"}
    assert p["lines"] == [{"no": 1, "description": "SAMPLE acrylic signboard 60x30", "qty": "40", "uom": "pcs", "unit_price": "60.0000",
                           "discount_pct": None, "amount": "2400.00", "is_header": False, "section": None}]
    assert (p["subtotal"], p["tax_total"], p["total"]) == ("2400.00", None, "2400.00")
    assert p["issuer"]["email"] == "me@shop.test" and p["issuer"]["brand_color"] == "#aa0000" and p["issuer"]["bank_details"] == "Bank 123"
    assert [(c["channel"], c["available"]) for c in p["channels"]] == [("email", True), ("whatsapp", True), ("telegram", False)]
    assert p["channels"][0]["to_hint"] == "b***@cafe.test" and "buyer" not in json.dumps(p["channels"]) and p["channels"][2]["reason"] == "no_contact"
    t = pdf_text(e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content)  # the same figures and wording as the PDF
    assert "2,400.00" in t and "60.00" in t and "Ms Lim" in t and "9 Jalan Kopi" in t and "1 Jalan Test" in t
    s = e.c.get(f"/api/documents/{q}/preview", headers=e.S)
    assert s.status_code == 200 and s.json() == p and not keys(s.json()) & STAFF_HIDE
    assert not any(w in json.dumps(p).lower() for w in ("cost", "margin", "supplier"))
    inbound = e.sql("select id from documents where direction='in' limit 1")[0]["id"]
    assert e.c.get(f"/api/documents/{inbound}/preview", headers=e.O).status_code == 404
    assert e.c.get(f"/api/documents/{q}/preview", headers=other_tenant(e)).status_code == 404


def test_preview_channels_read_the_contact_forms(env):
    e = env
    q = seed_quote(e)
    pid = "(select party_id from documents where id=%s)"
    ch = lambda: {c["channel"]: c for c in e.c.get(f"/api/documents/{q}/preview", headers=e.O).json()["channels"]}
    e.sql(f"update parties set contact='tg:4242' where id={pid}", (q,))
    c = ch()
    assert (c["telegram"]["available"], c["whatsapp"]["available"], c["whatsapp"]["reason"]) == (True, False, "no_contact")  # tg: is not a phone number
    assert c["telegram"]["to_hint"] == "***242" and c["email"]["reason"] == "no_email"
    e.sql(f"update parties set contact=null where id={pid}", (q,))
    assert not any(x["available"] for x in ch().values())


# ---------------------------------------------------------------- 4. exports
def test_customer_exports_in_every_format_carry_no_cost(env):
    e = env
    q = seed_quote(e)
    mimes = {"pdf": "application/pdf", "xlsx": "spreadsheetml", "csv": "text/csv"}
    for who in (e.O, e.S):
        for fmt, mime in mimes.items():
            r = export(e, q, fmt, who=who)
            assert r.status_code == 200 and mime in r.headers["content-type"], (fmt, r.status_code)
            assert r.headers["content-disposition"] == f'attachment; filename="SAMPLE-QT-0001.{fmt}"'
            text = (pdf_text(r.content) if fmt == "pdf" else " ".join(str(c) for row in xlsx_rows(r.content) for c in row if c is not None)
                    if fmt == "xlsx" else r.content.decode("utf-8-sig")).lower()
            assert "2400" in text.replace(",", "") and not any(w in text for w in COST_WORDS), (fmt, [w for w in COST_WORDS if w in text])
    assert e.c.get(f"/api/documents/{q}/export", headers=e.O).content[:5] == b"%PDF-"  # default = pdf / customer
    rows = xlsx_rows(export(e, q, "xlsx").content)
    assert rows[0][:2] == ("Quotation", "SAMPLE-QT-0001") and any(r[5] == "Total" and r[6] == 2400 for r in rows)
    head = next(r for r in rows if r[0] == "#")
    assert head[:3] == ("#", "Description", "Qty") and len(head) == 7  # no cost columns
    crow = csv_rows(export(e, q, "csv"))
    assert crow[0] == ["Quotation", "SAMPLE-QT-0001"] and ["1", "SAMPLE acrylic signboard 60x30", "40", "pcs", "60.0000", "", "2400.00"] in crow
    assert not e.sql("select 1 from audit_log where event like 'quote exported%%'")  # the customer copy is not an audited export


def test_internal_export_needs_both_permissions_and_is_audited(env):
    e = env
    q = seed_quote(e)
    only_export = member(e, role(e, "Exporter", ["export_data"]), "Ex")
    only_costs = member(e, role(e, "Coster", ["view_costs"]), "Co")
    both = member(e, role(e, "Both", ["export_data", "view_costs"]), "Bo")
    for who in (e.S, only_export, only_costs):
        for fmt in ("xlsx", "csv"):
            r = export(e, q, fmt, "internal", who)
            assert r.status_code == 403 and r.json()["error"]["code"] == "permission_required", (who, fmt)
    assert export(e, q, "pdf", "internal", e.O).status_code == 415 and export(e, q, "pdf", "internal", e.S).status_code == 403
    assert export(e, q, "xlsx", "internal", both).status_code == 200
    r = export(e, q, "xlsx", "internal", e.O)
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('-internal.xlsx"')
    rows = xlsx_rows(r.content)
    head = next(x for x in rows if x[0] == "#")
    assert head[7:] == ("Cost (MYR)", "Margin (MYR)", "Margin %", "Cost stale")
    line = next(x for x in rows if x[0] == 1)
    assert line[6:10] == (2400.0, 1500.0, 900.0, 37.5) and not line[10]
    by = {x[0]: x[1] for x in rows if x and x[0] in ("Total cost", "Total margin", "Margin %", "Any cost stale")}
    assert by == {"Total cost": 1500.0, "Total margin": 900.0, "Margin %": 37.5, "Any cost stale": "no"}
    c = csv_rows(export(e, q, "csv", "internal", e.O))
    assert ["1", "SAMPLE acrylic signboard 60x30", "40", "pcs", "60.0000", "", "2400.00", "1500.00", "900.00", "37.50", ""] in c
    e.shock()  # acrylic +40 %: the quote's cost basis is now stale
    assert dict((x[0], x[1]) for x in xlsx_rows(export(e, q, "xlsx", "internal").content) if x[0] == "Any cost stale")["Any cost stale"] == "yes"
    ev = [x["event"] for x in e.sql("select event from audit_log where entity='documents' and entity_id=%s", (q,))]
    assert ev.count("quote exported (internal xlsx)") == 3 and "quote exported (internal csv)" in ev
    d = draft(e)  # unlike the BOM export, a draft may be exported
    assert export(e, d, "xlsx", "internal").status_code == 200 and export(e, d, "csv", "internal").status_code == 200


def test_export_bad_params_and_scope(env):
    e = env
    q = seed_quote(e)
    for qs in ("format=docx", "variant=secret", "format=PDF", "format=&variant="):
        assert e.c.get(f"/api/documents/{q}/export?{qs}", headers=e.O).status_code == 400, qs
    assert e.c.get(f"/api/documents/{q}/export?format=zip", headers=e.O).json()["error"]["code"] == "bad_param"
    inbound = e.sql("select id from documents where direction='in' limit 1")[0]["id"]
    assert export(e, inbound).status_code == 404 and export(e, inbound, "xlsx", "internal").status_code == 404
    assert export(e, q, "xlsx", "customer", other_tenant(e)).status_code == 404
    assert export(e, q, "xlsx", "internal", other_tenant(e)).status_code == 404
    assert e.c.get(f"/api/documents/{q}/pdf", headers=e.O).content[:5] == b"%PDF-"  # the old route still works


def test_exports_neutralise_formulas_in_every_text_cell(env):
    e = env
    q = seed_quote(e)
    e.sql("update parties set name='=HYPERLINK(\"x\")', attention='+cmd', address='@sum', email='-a@b.co' where id=(select party_id from documents where id=%s)", (q,))
    e.sql("update document_lines set description='=1+1', uom='-x' where document_id=%s", (q,))
    e.sql("update documents set payment_terms='=pay', terms_text='@terms' where id=%s", (q,))
    e.sql("update tenants set address='+addr', bank_details='-bank', issuer_phone='=tel'")
    cells = [c for r in xlsx_rows(export(e, q, "xlsx").content) for c in r if isinstance(c, str)]
    cells += [c for r in csv_rows(export(e, q, "csv")) for c in r]
    assert not [c for c in cells if c and c[0] in "=+-@\t\r"], [c for c in cells if c and c[0] in "=+-@\t\r"]
    for want in ("'=HYPERLINK(\"x\")", "'+cmd", "'@sum", "'-a@b.co", "'=1+1", "'-x", "'=pay", "'@terms", "'+addr", "'-bank", "'=tel"):
        assert want in cells, want
    internal = [c for r in csv_rows(export(e, q, "csv", "internal")) for c in r]
    assert not [c for c in internal if c and c[0] in "=+@\t\r"] and "'=1+1" in internal  # negative numbers stay numbers


# ---------------------------------------------------------------- 5. sending
def test_email_send_attaches_the_pdf_moves_approved_to_sent_and_resends(env):
    e = env
    d = approved(e, email="buyer@cafe.test")
    r = send(e, d, "email", message="Thanks <b>again</b>")
    assert r.status_code == 200 and r.json() == {"delivery": {"channel": "email", "status": "delivered"}, "status": "sent"}
    to, subject, text, html, cat = e.svc.outbox[-1]
    att = e.svc.attached[-1]
    assert to == "buyer@cafe.test" and subject.startswith("Quotation QT-") and cat == "quote"
    assert len(att) == 1 and att[0]["filename"].endswith(".pdf") and att[0]["mime"] == "application/pdf" and att[0]["content"][:5] == b"%PDF-"
    assert "PDF" in text and "Valid until" in text and "Thanks <b>again</b>" in text and "&lt;b&gt;again&lt;/b&gt;" in html and "<b>again" not in html
    assert not any(w in (text + html + subject + pdf_text(att[0]["content"])).lower() for w in ("cost", "margin", "supplier"))
    sent_at = e.sql("select sent_at from documents where id=%s", (d,))[0]["sent_at"]
    assert sent_at is not None
    m = e.sql("select channel::text c, direction::text d, body, external_id from messages where document_id=%s and channel='email'", (d,))
    assert len(m) == 1 and m[0]["d"] == "out" and len(m[0]["body"]) < 200 and "%PDF" not in m[0]["body"]
    ev = [x["event"] for x in e.sql("select event from audit_log where entity='documents' and entity_id=%s", (d,))]
    assert "quote send by email: delivered" in ev and "status approved->sent" in ev
    assert "buyer@cafe.test" not in json.dumps(e.sql("select event, before, after from audit_log"), default=str)
    r = send(e, d, "email")  # resend: status and sent_at stay
    assert r.json()["status"] == "sent" and r.json()["delivery"]["status"] == "delivered" and len(e.svc.outbox) == 2
    assert e.sql("select sent_at from documents where id=%s", (d,))[0]["sent_at"] == sent_at
    tr(e, d, "read", e.S)
    assert send(e, d, "email").json()["status"] == "read"


def test_whatsapp_and_telegram_send_a_document(env):
    e = env
    q = seed_quote(e)  # sent, customer contact +60-000-0000009
    r = send(e, q, "whatsapp", message="Hello")
    assert r.json() == {"delivery": {"channel": "whatsapp", "status": "delivered"}, "status": "sent"}
    ch, to, name, caption, pdf, pn = e.svc.docs[-1]
    assert (ch, to, name) == ("whatsapp", "+60-000-0000009", "SAMPLE-QT-0001.pdf") and "SAMPLE-QT-0001" in caption and "Hello" in caption and pdf[:5] == b"%PDF-"
    e.sql("update parties set contact='tg:4242' where id=(select party_id from documents where id=%s)", (q,))
    assert send(e, q, "whatsapp").json()["error"]["code"] == "no_contact"  # a tg: contact is never a phone number
    assert send(e, q, "telegram").json()["delivery"]["status"] == "delivered" and e.svc.docs[-1][:2] == ("telegram", "4242")
    e.sql("update parties set contact='+60123456789' where id=(select party_id from documents where id=%s)", (q,))
    assert send(e, q, "telegram").json()["error"]["code"] == "no_contact"
    assert send(e, q, "email").json()["error"]["code"] == "no_email"
    assert send(e, q, "email").status_code == 422 and send(e, q, "sms").status_code == 422
    assert e.sql("select count(*) n from messages where document_id=%s and body like 'Quotation%% sent as a PDF%%'", (q,))[0]["n"] == 2
    assert e.sql("select status::text s from documents where id=%s", (q,))[0]["s"] == "sent"


def test_send_refuses_wrong_status_and_non_owner(env):
    e = env
    d = draft(e)
    e.sql("update parties set email='a@b.co' where id=(select party_id from documents where id=%s)", (d,))
    r = send(e, d)
    assert r.status_code == 409 and r.json()["error"]["code"] == "not_sendable" and not e.svc.outbox  # a draft is approved first
    tr(e, d, "approved")
    assert send(e, d, who=e.S).status_code == 403 and send(e, d, who=member(e, role(e, "Ex", ["export_data", "view_costs"]), "Ex")).status_code == 403
    assert not e.svc.outbox
    send(e, d)
    tr(e, d, "accepted")
    assert send(e, d).status_code == 409 and len(e.svc.outbox) == 1
    inbound = e.sql("select id from documents where direction='in' limit 1")[0]["id"]
    assert send(e, inbound).status_code == 404
    assert e.c.post(f"/api/documents/{d}/send", json={"channel": "email", "extra": 1}, headers=e.O).status_code == 422
    assert e.c.post(f"/api/documents/{d}/send", json={"channel": "email", "message": "x" * 501}, headers=e.O).status_code == 422


def test_failed_delivery_is_reported_and_an_approved_quote_stays_approved(env):
    e = env
    d = approved(e, email="a@b.co", contact="+60123456789")
    e.svc.mail = lambda *a, **k: False  # email off or refused
    r = send(e, d)
    assert r.status_code == 200 and r.json()["delivery"]["status"] == "failed" and "not accepted" in r.json()["delivery"]["error"]
    assert r.json()["status"] == "approved" and e.sql("select sent_at from documents where id=%s", (d,))[0]["sent_at"] is None
    e.svc.send_document = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("outside WhatsApp's 24h window and no template"))
    r = send(e, d, "whatsapp")
    assert r.json()["delivery"] == {"channel": "whatsapp", "status": "failed", "error": "outside WhatsApp's 24h window and no template"}
    assert r.json()["status"] == "approved"
    e.svc.send_document = lambda *a, **k: (_ for _ in ()).throw(httpx.ConnectError("https://secret.example/token"))
    r = send(e, d, "whatsapp")
    assert r.json()["delivery"]["status"] == "failed" and "secret" not in r.text
    assert e.sql("select count(*) n from messages where document_id=%s and direction='out'", (d,))[0]["n"] == 0  # nothing recorded as sent
    assert any(x["event"] == "quote send by whatsapp: failed" for x in e.sql("select event from audit_log"))
    e.svc.mail = lambda *a, **k: True
    assert send(e, d).json()["status"] == "sent"  # retry works


def test_send_is_throttled_per_business(env):
    e = env
    q = seed_quote(e)
    e.sql("update parties set email='a@b.co' where id=(select party_id from documents where id=%s)", (q,))
    for _ in range(20):
        assert send(e, q).status_code == 200
    r = send(e, q)
    assert r.status_code == 429 and len(e.svc.outbox) == 20
    assert send(e, q, "telegram").status_code == 422  # validation errors do not use up the allowance


def test_send_is_blocked_without_a_licence_and_isolated_per_tenant(env, monkeypatch):
    e = env
    q = seed_quote(e)
    e.sql("update parties set email='a@b.co' where id=(select party_id from documents where id=%s)", (q,))
    assert send(e, q, who=other_tenant(e)).status_code == 404  # another business cannot reach this quote
    monkeypatch.setenv("LICENSE_PUBLIC_KEY", PUB)
    Net().adopt(monkeypatch)
    lic.invalidate()
    try:
        r = send(e, q)
        assert r.status_code == 402 and r.json()["error"]["code"] == "license_required" and not e.svc.outbox
        assert export(e, q, "xlsx").status_code == 200 and e.c.get(f"/api/documents/{q}/preview", headers=e.O).status_code == 200  # reads stay open
    finally:
        lic.invalidate()


def test_approving_then_sending_over_the_enquiry_channel_adds_the_pdf(env):
    e = env
    d = draft(e, channel="telegram")
    e.sql("update parties set contact='4242' where id=(select party_id from documents where id=%s)", (d,))
    tr(e, d, "approved")
    n = len(e.svc.sent)
    assert tr(e, d, "sent").json()["delivery"] == "delivered"
    assert len(e.svc.sent) == n + 1 and e.svc.sent[-1][:2] == ("telegram", "4242")  # the text as before
    assert [x[:3] for x in e.svc.docs] == [("telegram", "4242", e.svc.docs[0][2])] and e.svc.docs[0][4][:5] == b"%PDF-"
    d2 = draft(e, channel="telegram")
    e.sql("update parties set contact='4242' where id=(select party_id from documents where id=%s)", (d2,))
    tr(e, d2, "approved")
    e.svc.send_document = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no"))  # PDF fails, the text already went: still delivered
    assert tr(e, d2, "sent").json()["delivery"] == "delivered"


def test_sending_with_no_chat_falls_back_to_email_and_says_not_sent_without_one(env):
    e = env
    d = approved(e)  # simulated web enquiry: no chat contact, no email
    assert tr(e, d, "sent").json()["delivery"] == "not_sent" and not e.svc.outbox
    d2 = draft(e)
    e.sql("update parties set email='buyer@cafe.test' where id=(select party_id from documents where id=%s)", (d2,))
    tr(e, d2, "approved")
    assert tr(e, d2, "sent").json()["delivery"] == "delivered"
    assert e.svc.outbox[-1][0] == "buyer@cafe.test" and e.svc.attached[-1][0]["content"][:5] == b"%PDF-"


# ---------------------------------------------------------------- channel layer and mailer
def test_mailer_attachments_are_base64_and_capped(monkeypatch):
    seen = []
    monkeypatch.setenv("MAILTRAP_API_TOKEN", "tok")
    monkeypatch.setattr(mailer.httpx, "post", lambda url, **k: (seen.append((url, k["json"])), httpx.Response(200, request=httpx.Request("POST", url)))[1])
    assert mailer.send("a@b.co", "s", "t", None, "quote", attachments=[{"filename": "q.pdf", "content": b"%PDF-1", "mime": "application/pdf"}])
    att = seen[0][1]["attachments"][0]
    assert att == {"filename": "q.pdf", "type": "application/pdf", "disposition": "attachment", "content": base64.b64encode(b"%PDF-1").decode()}
    assert mailer.send("a@b.co", "s", "t") and "attachments" not in seen[1][1]
    big = [{"filename": "q.pdf", "content": bytes(mailer.MAX_ATTACH + 1), "mime": "application/pdf"}]
    assert mailer.send("a@b.co", "s", "t", attachments=big) is False and len(seen) == 2  # refused before any request
    monkeypatch.setattr(mailer.httpx, "post", lambda *a, **k: (_ for _ in ()).throw(httpx.ConnectError("x")))
    assert mailer.send("a@b.co", "s", "t", attachments=[big[0] | {"content": b"x"}]) is False  # still never raises


def test_whatsapp_and_telegram_document_calls(monkeypatch):
    m = Meta()
    kw = dict(client=m.client, token="tok", phone_id="PID")
    assert whatsapp.send_document("6012", b"%PDF-x", "application/pdf", "q.pdf", "Quote", **kw).startswith("wamid.out")
    assert m.calls[0][0] == "media" and b"%PDF-x" in m.calls[0][2]
    assert m.calls[1][2]["type"] == "document" and m.calls[1][2]["document"] == {"id": "MEDIA1", "filename": "q.pdf", "caption": "Quote"}
    seen = []
    c = httpx.Client(transport=httpx.MockTransport(lambda r: (seen.append(r), httpx.Response(200, json={"result": {"message_id": 9, "chat": {"id": 4242}}}))[1]))
    assert telegram.send_document(4242, b"%PDF-x", "application/pdf", "q.pdf", "x" * 2000, client=c, token="T") == "4242:9"
    assert seen[0].url.path == "/botT/sendDocument" and b"%PDF-x" in seen[0].content and b'filename="q.pdf"' in seen[0].content
    assert b"x" * 1024 in seen[0].content and b"x" * 1025 not in seen[0].content  # caption cap
    nope = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(403, text="blocked")))
    with pytest.raises(RuntimeError) as ei:
        telegram.send_document(4242, b"x", "application/pdf", "q.pdf", "c", client=nope, token="SECRET")
    assert "SECRET" not in str(ei.value) and "Start" in str(ei.value)


def test_document_outside_the_24h_window_needs_the_followup_template(monkeypatch):
    monkeypatch.setenv("WHATSAPP_TOKEN", "t")
    monkeypatch.setenv("WHATSAPP_PHONE_NUMBER_ID", "1")
    monkeypatch.delenv("WHATSAPP_FOLLOWUP_TEMPLATE", raising=False)
    mm = Meta(fail={"6012": (("document",), 131047)})
    s = Services()
    s.http = mm.client
    with pytest.raises(RuntimeError, match="WHATSAPP_FOLLOWUP_TEMPLATE"):
        s.send_document("whatsapp", "6012", b"%PDF-x", "q.pdf", "Quote QT-1")
    monkeypatch.setenv("WHATSAPP_FOLLOWUP_TEMPLATE", "quote_followup")
    assert s.send_document("whatsapp", "6012", b"%PDF-x", "q.pdf", "Quote  QT-1").endswith(".6012.template")
    assert mm.calls[-1][2]["template"]["components"][0]["parameters"][0]["text"] == "Quote QT-1"
    s.http = Meta(fail={"6012": (("document",), 100)}).client  # any other refusal is not hidden by the template
    with pytest.raises(httpx.HTTPStatusError):
        s.send_document("whatsapp", "6012", b"%PDF-x", "q.pdf", "Quote")


# ---------------------------------------------------------------- schema
def test_channel_enum_gains_email_once_and_reapplying_is_harmless(env, pg_url):
    for _ in range(2):
        db.apply_schema(seed=True)
    with psycopg.connect(pg_url, autocommit=True) as c:
        labels = [x[0] for x in c.execute("select unnest(enum_range(null::channel_t))::text")]
        cols = {x[0] for x in c.execute("select column_name from information_schema.columns where table_name='tenants' and (column_name like 'quote_%' or column_name like 'issuer_%')")}
    assert labels.count("email") == 1 and labels[:4] == ["telegram", "whatsapp", "web", "upload"]
    assert cols == {"quote_validity_days", "quote_payment_terms", "quote_terms_text", "quote_prefix", "issuer_email", "issuer_phone"}
    with pytest.raises(psycopg.errors.CheckViolation):
        env.sql("update tenants set quote_validity_days=0")
    with pytest.raises(psycopg.errors.CheckViolation):
        env.sql("update tenants set quote_prefix='lower'")
